#!/usr/bin/env python3
"""
Scan email for job-application-related messages and export to Excel.
Optional: .ics reminders for positive signals; optional macOS Reminders.

Gmail: place OAuth client JSON as credentials.json in this folder, run once to authorize.

IMAP: set IMAP_HOST, IMAP_USER, IMAP_PASSWORD in .env (see .env.example).

Usage:
  python sync.py --source gmail --out output/applications.xlsx
  python sync.py --source gmail --out output/applications.xlsx --ics output/reminders.ics --mac-reminders
"""

from __future__ import annotations

import argparse
import base64
import email as email_lib
import html
import imaplib
import json
import os
import re
import subprocess
import sys
import time
import uuid
from dataclasses import dataclass, field
from datetime import datetime, timedelta, timezone
from email.header import decode_header, make_header
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
from urllib.request import Request as UrlRequest, urlopen
import hashlib
from openpyxl import Workbook
from openpyxl.styles import Font

# Gmail (optional import until used)
try:
    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build
except ImportError:
    Credentials = None  # type: ignore

SCOPES = ["https://www.googleapis.com/auth/gmail.readonly"]


@dataclass
class ParsedEmail:
    message_id: str
    thread_id: str
    date: datetime
    subject: str
    from_header: str
    snippet: str
    body_text: str
    source: str = "gmail"


def sha256_text(s: str) -> str:
    return hashlib.sha256((s or "").encode("utf-8", errors="ignore")).hexdigest()


def read_json(path: Path, default: Any) -> Any:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def write_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, indent=2) + "\n", encoding="utf-8")


def load_config(path: Path) -> dict[str, Any]:
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def decode_mime_header(value: str) -> str:
    if not value:
        return ""
    try:
        return str(make_header(decode_header(value)))
    except Exception:
        return value


def extract_urls(text: str) -> list[str]:
    if not text:
        return []
    pattern = r"https?://[^\s<>\"')\]]+"
    return re.findall(pattern, text)


def extract_email_address(from_header: str) -> str:
    if not from_header:
        return ""
    m = re.search(r"<([^>]+)>", from_header)
    if m:
        return m.group(1).strip().lower()
    m = re.search(r"([A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,})", from_header, re.I)
    return (m.group(1).strip().lower() if m else "")


def clean_sender_name(from_header: str) -> str:
    m = re.match(r"^([^<]+)<", from_header)
    if not m:
        return ""
    name = m.group(1).strip().strip('"').strip()
    name = re.sub(r"\s*\((do not reply|no reply|noreply)\)\s*$", "", name, flags=re.I)
    name = re.sub(r"\s*@\s*(icims|workday)\s*$", "", name, flags=re.I)
    return name.strip()[:200]


def extract_company(subject: str, from_header: str) -> str:
    s = decode_mime_header(subject).strip()
    s2 = re.sub(r"^(re|fwd?):\s*", "", s, flags=re.I).strip()

    patterns = [
        r"(?i)^update on\s+(.+?)\s+application\b",
        r"(?i)^information about your application to\s+(.+?)\b",
        r"(?i)^your application to\s+(.+?)(?:\s*[-–|].*|\s*$)",
        r"(?i)^thank you for applying (?:to|at)\s+(.+?)(?:[.,]|$)",
        r"(?i)^thank you for your application to\s+(.+?)(?:[.,]|$)",
        r"(?i)^application\s+(?:for|to)\s+(.+?)(?:\s*[-–|].*|\s*$)",
        r"(?i)applying\s+to\s+(.+?)(?:\s*[-–|].*|\s*$)",
        r"(?i)^(.+?)\s+application\s+(?:update|status)\b",
        r"(?i)^(.+?)\s*[-–|]\s*application\s+(?:update|status)\b",
    ]
    for p in patterns:
        m = re.search(p, s2)
        if m:
            cand = m.group(1).strip(" -–|,")[:200]
            if cand and not re.fullmatch(r"(?i)(your|application|update|status)", cand):
                return cand

    email_addr = extract_email_address(from_header)
    domain = email_addr.split("@", 1)[1] if "@" in email_addr else ""
    sender_name = clean_sender_name(from_header)

    ats_domains = (
        "talent.icims.com",
        "myworkday.com",
        "hire.lever.co",
        "greenhouse.io",
        "mail.greenhouse.io",
        "smartrecruiters.com",
        "jobs.ashbyhq.com",
        "ashbyhq.com",
        "workablemail.com",
        "successfactors.com",
        "taleo.net",
        "oraclecloud.com",
        "hirebridge.com",
        "hirebridgemail.com",
    )
    if domain and any(domain.endswith(d) for d in ats_domains):
        if sender_name:
            return sender_name

    if domain and domain not in ("gmail.com", "googlemail.com", "outlook.com", "hotmail.com"):
        base = domain.split(".")[0].replace("-", " ").title()
        if base and base.lower() not in ("mail", "email", "noreply", "notification", "notifications"):
            return base[:200]

    if sender_name and "noreply" not in sender_name.lower():
        return sender_name[:200]

    return "(unknown)"


def normalize_for_match(text: str) -> str:
    t = html.unescape(text or "")
    return re.sub(r"\s+", " ", t.lower())


def extract_domain(url: str) -> str:
    try:
        return urlparse(url).netloc.lower()
    except Exception:
        return ""


def find_signal_hits(text: str, urls: list[str], config: dict[str, Any]) -> tuple[list[str], list[str]]:
    """
    Returns (strong_hits, soft_hits).
    Strong hits should be treated as Action needed.
    Soft hits should be Review soon (to reduce false positives).
    """
    t = normalize_for_match(text)
    domains = [extract_domain(u) for u in urls]

    strong: list[str] = []
    soft: list[str] = []

    # Link domains: strong, because they usually point to an OA, scheduler, or meeting.
    for dd in config.get("strong_action_link_domains", []):
        d = (dd or "").strip().lower()
        if not d:
            continue
        if any(dom == d or dom.endswith("." + d) for dom in domains):
            strong.append(f"link:{d}")

    # Keyword hits (strong/soft lists).
    for kw in config.get("strong_action_keywords", []):
        nkw = normalize_for_match(kw)
        if nkw and nkw in t:
            strong.append(kw)
    for kw in config.get("soft_action_keywords", []):
        nkw = normalize_for_match(kw)
        if nkw and nkw in t:
            soft.append(kw)

    # Heuristic escalation: interview + scheduling language => strong.
    if "interview" in t:
        if any(x in t for x in ("schedule", "scheduled", "reschedule", "rescheduled", "availability", "confirm")):
            strong.append("interview+scheduling")

    # De-duplicate while preserving order.
    def dedup(xs: list[str]) -> list[str]:
        out: list[str] = []
        seen: set[str] = set()
        for x in xs:
            if x not in seen:
                out.append(x)
                seen.add(x)
        return out

    strong = dedup(strong)
    soft = dedup([x for x in soft if x not in set(strong)])
    return strong, soft


def classify(text: str, urls: list[str], config: dict[str, Any]) -> tuple[str, list[str], str]:
    """
    Returns (status, hits, strength)
      - strength: 'strong' | 'soft' | ''
    """
    t = normalize_for_match(text)
    strong_hits, soft_hits = find_signal_hits(text, urls, config)

    # Strong opportunity signals should always be Action needed (high recall for OAs/interviews).
    if strong_hits:
        return "Action needed", strong_hits, "strong"

    # Rejections should override soft signals to avoid polluting review lists.
    for kw in config.get("rejection_keywords", []):
        if normalize_for_match(kw) in t:
            return "Rejection", [], ""

    # Soft signals are worth reviewing, but shouldn't trigger reminders.
    if soft_hits:
        return "Review soon", soft_hits, "soft"
    for kw in config.get("positive_keywords", []):
        if normalize_for_match(kw) in t:
            return "Action needed", [], ""
    for kw in config.get("confirmation_keywords", []):
        if normalize_for_match(kw) in t:
            return "Applied (confirmation)", [], ""
    return "Other / review", [], ""


def ollama_available(base_url: str) -> bool:
    try:
        req = UrlRequest(f"{base_url.rstrip('/')}/api/tags", headers={"Accept": "application/json"})
        with urlopen(req, timeout=1) as resp:
            return 200 <= resp.status < 300
    except Exception:
        return False


def ollama_parse_email(
    base_url: str,
    model: str,
    subject: str,
    from_header: str,
    snippet: str,
    body_text: str,
    urls: list[str],
    timeout_s: int = 60,
) -> dict[str, Any]:
    """
    Local, free parsing via Ollama. Returns a dict with keys:
      action_required(bool), action_type(str), company(str), role(str),
      deadline_utc(str|''), summary(str), confidence(float 0..1)
    """
    prompt = f"""
You are classifying recruiting emails for a job-application tracker.

Return ONLY valid JSON, no markdown, no extra keys.

JSON schema:
{{
  "action_required": boolean,   // true if user must take an action (OA/interview scheduling/complete assessment/respond)
  "action_type": string,        // one of: "OA", "Interview", "Scheduler", "RecruiterReplyNeeded", "Offer", "Rejection", "Confirmation", "Other"
  "company": string,
  "role": string,
  "deadline_utc": string,       // ISO-like "YYYY-MM-DD HH:MM" in UTC if present, else ""
  "summary": string,            // <= 200 chars
  "confidence": number          // 0.0 to 1.0
}}

Heuristics:
- If the email is a rejection, set action_required=false and action_type="Rejection".
- If it's just application received/thank you, action_required=false and action_type="Confirmation".
- If it contains an assessment/interview invitation, scheduling request, or link to complete something, action_required=true.

Email:
Subject: {subject}
From: {from_header}
Snippet: {snippet}
URLs: {", ".join(urls[:6])}
Body (truncated): {body_text[:4000]}
""".strip()

    payload = {
        "model": model,
        "prompt": prompt,
        "stream": False,
        "format": "json",
        "options": {"temperature": 0.1},
    }
    req = UrlRequest(
        f"{base_url.rstrip('/')}/api/generate",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json", "Accept": "application/json"},
        method="POST",
    )
    with urlopen(req, timeout=timeout_s) as resp:
        raw = resp.read().decode("utf-8", errors="replace")
    data = json.loads(raw)
    txt = data.get("response", "") if isinstance(data, dict) else ""
    out = json.loads(txt) if isinstance(txt, str) else {}
    if not isinstance(out, dict):
        return {}
    # Minimal normalization/safety
    out.setdefault("action_required", False)
    out.setdefault("action_type", "Other")
    out.setdefault("company", "")
    out.setdefault("role", "")
    out.setdefault("deadline_utc", "")
    out.setdefault("summary", "")
    out.setdefault("confidence", 0.0)
    return out


def gmail_body_from_payload(payload: dict[str, Any]) -> str:
    parts: list[str] = []

    def walk(p: dict[str, Any]) -> None:
        mime = p.get("mimeType", "")
        body = p.get("body", {})
        data = body.get("data")
        if data and mime == "text/plain":
            try:
                raw = base64.urlsafe_b64decode(data + "===")
                parts.append(raw.decode("utf-8", errors="replace"))
            except Exception:
                pass
        for sub in p.get("parts", []) or []:
            walk(sub)

    walk(payload)
    return "\n".join(parts)


def gmail_list_and_parse(
    service: Any, query: str, config: dict[str, Any], max_results: int
) -> list[ParsedEmail]:
    out: list[ParsedEmail] = []
    page_token = None
    fetched = 0
    while fetched < max_results:
        batch = min(100, max_results - fetched)
        req = (
            service.users()
            .messages()
            .list(userId="me", q=query, maxResults=batch, pageToken=page_token)
        )
        resp = req.execute()
        msgs = resp.get("messages", [])
        if not msgs:
            break
        for m in msgs:
            if fetched >= max_results:
                break
            mid = m["id"]
            full = (
                service.users()
                .messages()
                .get(userId="me", id=mid, format="full")
                .execute()
            )
            headers = {h["name"].lower(): h["value"] for h in full.get("payload", {}).get("headers", [])}
            subject = headers.get("subject", "")
            from_h = headers.get("from", "")
            date_hdr = headers.get("date", "")
            try:
                dt = email_lib.utils.parsedate_to_datetime(date_hdr)
                if dt.tzinfo is None:
                    dt = dt.replace(tzinfo=timezone.utc)
            except Exception:
                dt = datetime.now(timezone.utc)
            internal = int(full.get("internalDate", "0")) / 1000.0
            if internal:
                dt = datetime.fromtimestamp(internal, tz=timezone.utc)
            snippet = full.get("snippet", "") or ""
            body = gmail_body_from_payload(full.get("payload", {}))
            if not body.strip():
                body = snippet
            out.append(
                ParsedEmail(
                    message_id=mid,
                    thread_id=full.get("threadId", ""),
                    date=dt,
                    subject=subject,
                    from_header=from_h,
                    snippet=snippet,
                    body_text=body[:12000],
                    source="gmail",
                )
            )
            fetched += 1
        page_token = resp.get("nextPageToken")
        if not page_token:
            break
    return out


def get_gmail_service(credentials_path: Path, token_path: Path) -> Any:
    if Credentials is None:
        raise SystemExit("Install google packages: pip install -r requirements.txt")
    creds = None
    if token_path.exists():
        creds = Credentials.from_authorized_user_file(str(token_path), SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            if not credentials_path.exists():
                raise SystemExit(
                    f"Missing {credentials_path}. Download OAuth client JSON from Google Cloud Console."
                )
            flow = InstalledAppFlow.from_client_secrets_file(str(credentials_path), SCOPES)
            creds = flow.run_local_server(port=0)
        token_path.write_text(creds.to_json(), encoding="utf-8")
    return build("gmail", "v1", credentials=creds, cache_discovery=False)


def imap_parse_messages(
    host: str,
    user: str,
    password: str,
    folder: str,
    max_results: int,
) -> list[ParsedEmail]:
    out: list[ParsedEmail] = []
    M = imaplib.IMAP4_SSL(host)
    try:
        M.login(user, password)
        M.select(folder)
        typ, data = M.search(None, "ALL")
        if typ != "OK" or not data[0]:
            return out
        ids = data[0].split()
        ids = ids[-max_results:] if len(ids) > max_results else ids
        for num in reversed(ids):
            typ, msg_data = M.fetch(num, "(RFC822)")
            if typ != "OK" or not msg_data or not isinstance(msg_data[0], tuple):
                continue
            raw = msg_data[0][1]
            msg = email_lib.message_from_bytes(raw)
            subj = decode_mime_header(msg.get("Subject", "") or "")
            from_h = decode_mime_header(msg.get("From", "") or "")
            date_hdr = msg.get("Date", "")
            try:
                dt = email_lib.utils.parsedate_to_datetime(date_hdr)
                if dt.tzinfo is None:
                    dt = dt.replace(tzinfo=timezone.utc)
            except Exception:
                dt = datetime.now(timezone.utc)
            body_chunks: list[str] = []
            if msg.is_multipart():
                for part in msg.walk():
                    if part.get_content_type() == "text/plain":
                        try:
                            body_chunks.append(
                                part.get_payload(decode=True).decode("utf-8", errors="replace")
                            )
                        except Exception:
                            pass
            else:
                try:
                    body_chunks.append(
                        msg.get_payload(decode=True).decode("utf-8", errors="replace")
                    )
                except Exception:
                    pass
            body = "\n".join(body_chunks)[:12000]
            mid = msg.get("Message-ID", "") or f"imap-{num.decode()}"
            snippet = (body or subj)[:280]
            out.append(
                ParsedEmail(
                    message_id=mid,
                    thread_id="",
                    date=dt,
                    subject=subj,
                    from_header=from_h,
                    snippet=snippet,
                    body_text=body,
                    source="imap",
                )
            )
    finally:
        try:
            M.logout()
        except Exception:
            pass
    return out


def parse_include_after(config: dict[str, Any]) -> datetime | None:
    raw = config.get("include_email_after")
    if not raw or not isinstance(raw, str):
        return None
    try:
        d = datetime.strptime(raw.strip(), "%Y-%m-%d")
        return d.replace(tzinfo=timezone.utc)
    except ValueError:
        return None


def write_excel(
    rows: list[dict[str, Any]],
    path: Path,
    meta: dict[str, Any] | None = None,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Applications"
    headers = [
        "Date (UTC)",
        "Company",
        "Status",
        "Signal strength",
        "Opportunity hits",
        "AI action_type",
        "AI action_required",
        "Role",
        "AI summary",
        "AI confidence",
        "Subject",
        "From",
        "First link",
        "Snippet",
        "Thread ID",
        "Message ID",
        "Source",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = Font(bold=True)
    for i, row in enumerate(rows, 2):
        for col, key in enumerate(
            [
                "date",
                "company",
                "status",
                "signal_strength",
                "opportunity_hits",
                "ai_action_type",
                "ai_action_required",
                "role",
                "ai_summary",
                "ai_confidence",
                "subject",
                "from",
                "link",
                "snippet",
                "thread_id",
                "message_id",
                "source",
            ],
            1,
        ):
            v = row.get(key, "")
            if key == "date" and isinstance(v, datetime):
                v = v.strftime("%Y-%m-%d %H:%M:%S")
            ws.cell(row=i, column=col, value=v)
    if meta:
        info = wb.create_sheet("Export info")
        info["A1"] = "Field"
        info["B1"] = "Value"
        info["A1"].font = Font(bold=True)
        info["B1"].font = Font(bold=True)
        exported = meta.get("exported_at_utc", "")
        lines = [
            ("Exported at (UTC)", exported),
            ("Include email on or after", meta.get("include_email_after", "")),
            ("Gmail search query", meta.get("gmail_query", "")),
            ("Total rows in Applications sheet", meta.get("total_rows", len(rows))),
            ("Action needed count", meta.get("action_needed", "")),
            ("Rejection count", meta.get("rejection", "")),
            ("Applied (confirmation) count", meta.get("applied", "")),
            ("Other (included) count", meta.get("other", "")),
        ]
        for r, (k, v) in enumerate(lines, 2):
            info.cell(row=r, column=1, value=k)
            info.cell(row=r, column=2, value=v)
        info.column_dimensions["A"].width = 28
        info.column_dimensions["B"].width = 72
    wb.save(path)


def write_ics(events: list[dict[str, Any]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        "BEGIN:VCALENDAR",
        "VERSION:2.0",
        "PRODID:-//job-tracker//EN",
        "CALSCALE:GREGORIAN",
    ]
    now = datetime.now(timezone.utc)
    for ev in events:
        uid = str(uuid.uuid4())
        start = ev.get("start", now + timedelta(hours=1))
        if start.tzinfo is None:
            start = start.replace(tzinfo=timezone.utc)
        end = start + timedelta(hours=1)
        fmt = "%Y%m%dT%H%M%SZ"

        def fmt_dt(d: datetime) -> str:
            return d.astimezone(timezone.utc).strftime(fmt)

        lines.extend(
            [
                "BEGIN:VEVENT",
                f"UID:{uid}",
                f"DTSTAMP:{fmt_dt(now)}",
                f"DTSTART:{fmt_dt(start)}",
                f"DTEND:{fmt_dt(end)}",
                f"SUMMARY:{ics_escape(ev.get('summary', 'Job action'))}",
                f"DESCRIPTION:{ics_escape(ev.get('description', ''))}",
                "END:VEVENT",
            ]
        )
    lines.append("END:VCALENDAR")
    path.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8")


def ics_escape(s: str) -> str:
    return (
        s.replace("\\", "\\\\")
        .replace(";", "\\;")
        .replace(",", "\\,")
        .replace("\n", "\\n")
    )[:8000]


def mac_reminder(title: str, when: datetime) -> None:
    if sys.platform != "darwin":
        return
    # Apple Reminders — default list
    iso = when.strftime("%Y-%m-%d %H:%M:%S")
    script = f'''
    tell application "Reminders"
        tell list "Reminders"
            make new reminder with properties {{name:{json.dumps(title)}, due date:date "{iso}"}}
        end tell
    end tell
    '''
    subprocess.run(["osascript", "-e", script], check=False, capture_output=True)


def load_dotenv(path: Path) -> None:
    if not path.exists():
        return
    for line in path.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#"):
            continue
        if "=" in line:
            k, _, v = line.partition("=")
            os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def main() -> None:
    root = Path(__file__).resolve().parent
    load_dotenv(root / ".env")

    parser = argparse.ArgumentParser(description="Export job-related email to Excel")
    parser.add_argument("--config", type=Path, default=root / "config.json")
    parser.add_argument("--source", choices=("gmail", "imap"), default="gmail")
    parser.add_argument("--out", type=Path, default=root / "output" / "applications.xlsx")
    parser.add_argument("--ics", type=Path, default=None, help="Write calendar reminders for Action needed")
    parser.add_argument("--mac-reminders", action="store_true", help="Create macOS Reminders (Action needed)")
    parser.add_argument(
        "--ai",
        choices=("none", "ollama"),
        default="ollama",
        help="Use local AI parsing (default: ollama). Falls back to none if unavailable.",
    )
    parser.add_argument(
        "--max",
        type=int,
        default=1000,
        dest="max_results",
        help="Max messages to fetch (default 1000)",
    )
    parser.add_argument(
        "--include-unmatched",
        action="store_true",
        help="Include rows classified as Other / review (default: only matched keywords)",
    )
    args = parser.parse_args()

    config = load_config(args.config)
    emails: list[ParsedEmail] = []
    output_dir = (args.out.parent if args.out else (root / "output"))
    ai_cache_path = output_dir / "ai-cache.json"
    ai_cache: dict[str, Any] = read_json(ai_cache_path, {})
    if not isinstance(ai_cache, dict):
        ai_cache = {}
    ai_base_url = str(config.get("ollama_url", "http://localhost:11434"))
    ai_model = str(config.get("ollama_model", "qwen2.5:7b-instruct"))
    ai_enabled = args.ai == "ollama" and ollama_available(ai_base_url)
    if args.ai == "ollama" and not ai_enabled:
        print("Ollama not available; continuing without AI. Start Ollama and pull a model, then re-run.")
    ai_parse_statuses = config.get("ai_parse_statuses", ["Action needed", "Review soon"])
    if not isinstance(ai_parse_statuses, list):
        ai_parse_statuses = ["Action needed", "Review soon"]
    ai_max_per_run = config.get("ai_max_emails_per_run", 150)
    try:
        ai_max_per_run = int(ai_max_per_run)
    except Exception:
        ai_max_per_run = 150
    ai_used = 0

    if args.source == "gmail":
        creds = root / "credentials.json"
        token = root / "token.json"
        service = get_gmail_service(creds, token)
        try:
            profile = service.users().getProfile(userId="me").execute()
            user_email = (profile.get("emailAddress") or "").strip().lower()
        except Exception:
            user_email = ""
        q = config.get("gmail_query", "")
        emails = gmail_list_and_parse(service, q, config, args.max_results)
    else:
        host = os.environ.get("IMAP_HOST", "")
        user = os.environ.get("IMAP_USER", "")
        password = os.environ.get("IMAP_PASSWORD", "")
        if not (host and user and password):
            raise SystemExit("Set IMAP_HOST, IMAP_USER, IMAP_PASSWORD in .env")
        folder = config.get("imap_folder", "INBOX")
        emails = imap_parse_messages(host, user, password, folder, args.max_results)

    cutoff = parse_include_after(config)
    if cutoff is not None:
        before = len(emails)
        emails = [e for e in emails if e.date >= cutoff]
        if before != len(emails):
            print(f"Date filter ({config.get('include_email_after')}): kept {len(emails)} of {before} messages.")

    rows: list[dict[str, Any]] = []
    ics_events: list[dict[str, Any]] = []

    for pe in emails:
        if args.source == "gmail" and user_email:
            if extract_email_address(pe.from_header) == user_email:
                continue
        blob = f"{pe.subject}\n{pe.snippet}\n{pe.body_text}"
        urls = extract_urls(pe.body_text or pe.snippet)
        status, opp_hits, strength = classify(blob, urls, config)
        company = extract_company(pe.subject, pe.from_header)
        link = urls[0] if urls else ""

        ai = {}
        if ai_enabled and status in ai_parse_statuses and ai_used < ai_max_per_run:
            cache_key = f"{pe.message_id}:{sha256_text(pe.subject + pe.snippet + pe.body_text[:4000])}"
            ai = ai_cache.get(cache_key) if isinstance(ai_cache.get(cache_key), dict) else None
            if not ai:
                # light throttling to keep the machine responsive
                time.sleep(float(config.get("ollama_min_delay_s", 0.0) or 0.0))
                try:
                    ai = ollama_parse_email(
                        ai_base_url,
                        ai_model,
                        pe.subject,
                        pe.from_header,
                        pe.snippet,
                        pe.body_text,
                        urls,
                        timeout_s=int(config.get("ollama_timeout_s", 60) or 60),
                    )
                except Exception:
                    ai = {}
                ai_cache[cache_key] = ai or {}
            ai_used += 1

            # Use AI as an additional safety net: if it says action_required, escalate.
            if ai.get("action_required") is True and status not in ("Rejection",):
                status = "Action needed"
                strength = strength or "ai"
                if not opp_hits:
                    opp_hits = ["ai:action_required"]

            # Prefer AI company/role when present (but keep rule-based fallback).
            ai_company = (ai.get("company") or "").strip()
            if ai_company and len(ai_company) >= 2:
                company = ai_company[:200]

        role = ""
        if isinstance(ai, dict):
            role = (ai.get("role") or "").strip()[:200]

        row = {
            "date": pe.date,
            "company": company,
            "status": status,
            "signal_strength": strength,
            "opportunity_hits": ", ".join(opp_hits)[:500],
            "ai_action_type": (ai.get("action_type") if isinstance(ai, dict) else "") or "",
            "ai_action_required": str(bool(ai.get("action_required"))) if isinstance(ai, dict) else "",
            "role": role,
            "ai_summary": (ai.get("summary") if isinstance(ai, dict) else "") or "",
            "ai_confidence": (ai.get("confidence") if isinstance(ai, dict) else ""),
            "subject": pe.subject,
            "from": pe.from_header,
            "link": link,
            "snippet": pe.snippet[:500],
            "thread_id": pe.thread_id,
            "message_id": pe.message_id,
            "source": pe.source,
        }
        if status == "Other / review" and not args.include_unmatched:
            continue
        rows.append(row)

        if status == "Action needed":
            start = pe.date.astimezone(timezone.utc) + timedelta(hours=2)
            title = f"Job: {company} — follow up / assessment"
            desc = f"{pe.subject}\n{link}"
            ics_events.append(
                {"start": start, "summary": title, "description": desc}
            )
            if args.mac_reminders:
                mac_reminder(title, datetime.now() + timedelta(minutes=30))

    rows.sort(key=lambda r: r["date"], reverse=True)
    exported_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    meta = {
        "exported_at_utc": exported_at,
        "include_email_after": config.get("include_email_after", ""),
        "gmail_query": config.get("gmail_query", "") if args.source == "gmail" else "(IMAP)",
        "total_rows": len(rows),
        "action_needed": sum(1 for r in rows if r["status"] == "Action needed"),
        "rejection": sum(1 for r in rows if r["status"] == "Rejection"),
        "applied": sum(1 for r in rows if r["status"] == "Applied (confirmation)"),
        "other": sum(1 for r in rows if r["status"] == "Other / review"),
    }
    write_excel(rows, args.out, meta=meta)
    print(f"Wrote {len(rows)} rows to {args.out}")
    if ai_enabled:
        write_json(ai_cache_path, ai_cache)

    if args.ics and ics_events:
        write_ics(ics_events, args.ics)
        print(f"Wrote {len(ics_events)} events to {args.ics} (open in Calendar or import)")

    action = sum(1 for r in rows if r["status"] == "Action needed")
    print(f"Status summary: Action needed={action}, others classified by keywords (tune config.json).")


if __name__ == "__main__":
    main()
